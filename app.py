import os, secrets
from pathlib import Path
from flask import Flask, render_template, request, jsonify
from openpyxl import load_workbook
from print_queue import init_db, enqueue, claim, finish, get_job
BASE=Path(__file__).resolve().parent; DATA=BASE/'data'/'master.xlsx'
app=Flask(__name__); BRIDGE_TOKEN=os.getenv('BRIDGE_TOKEN','CHANGE-ME-NOW'); rows=[]; init_db()
def text(v):
    if v is None:return ''
    if isinstance(v,float) and v.is_integer():return str(int(v))
    return str(v).strip()
def load_master():
    global rows
    wb=load_workbook(DATA,read_only=True,data_only=True);ws=wb[wb.sheetnames[0]];it=ws.iter_rows(values_only=True);h=[text(x).upper() for x in next(it)]
    names={'pid':['ITEM CODE','PRODUCT ID'],'name':['ITEM NAME','PRODUCT NAME'],'ean':['EAN CODE','EAN','BARCODE'],'uom':['CONVERSION','UOM']};idx={k:next((h.index(n) for n in a if n in h),None) for k,a in names.items()}
    if any(idx[k] is None for k in ('pid','name','ean')):raise ValueError('Required columns: ITEM CODE, ITEM NAME, EAN CODE')
    rows=[]
    for r in it:
        p=text(r[idx['pid']]);n=text(r[idx['name']]);e=text(r[idx['ean']]);u=text(r[idx['uom']]) if idx['uom'] is not None else ''
        if p or n or e:rows.append({'product_id':p,'product_name':n,'ean':e,'uom':u})
    wb.close()
load_master()
@app.get('/')
def home():return render_template('index.html')
@app.get('/api/search')
def search():
    q=text(request.args.get('q')).lower();exact=[];partial=[]
    for x in rows:
        v=(x['product_id'].lower(),x['ean'].lower(),x['product_name'].lower(),x['uom'].lower())
        if q in v[:2]:exact.append(x)
        elif q and any(q in z for z in v):partial.append(x)
    return jsonify((exact+partial)[:50])
@app.get('/api/stats')
def stats():return jsonify(records=len(rows),products=len({x['product_id'] for x in rows}))
@app.post('/api/print-jobs')
def create_job():
    d=request.get_json(silent=True) or {};z=str(d.get('zpl',''));bid=str(d.get('bridge_id','')).strip()[:80]
    try:q=max(1,min(999,int(d.get('qty',1))))
    except:return jsonify(error='Invalid quantity'),400
    if not bid or not z.startswith('^XA') or not z.endswith('^XZ'):return jsonify(error='Valid Bridge ID and ZPL required'),400
    jid=enqueue(bid,z[:50000],str(d.get('product_id',''))[:100],str(d.get('ean',''))[:100],q);return jsonify(ok=True,job_id=jid,status='pending')
@app.get('/api/print-jobs/<jid>')
def job_status(jid):
    j=get_job(jid);return (jsonify(j),200) if j else (jsonify(error='Not found'),404)
def auth():return secrets.compare_digest(request.headers.get('X-Bridge-Token',''),BRIDGE_TOKEN)
@app.post('/api/bridge/claim')
def bridge_claim():
    if not auth():return jsonify(error='Unauthorized'),401
    return jsonify(job=claim(str((request.get_json(silent=True) or {}).get('bridge_id',''))[:80]))
@app.post('/api/bridge/complete')
def bridge_complete():
    if not auth():return jsonify(error='Unauthorized'),401
    d=request.get_json(silent=True) or {};ok=finish(str(d.get('job_id','')),str(d.get('bridge_id','')),bool(d.get('ok')),str(d.get('message','')));return jsonify(ok=ok),(200 if ok else 409)
@app.get('/health')
def health():return jsonify(ok=True,records=len(rows))
if __name__=='__main__':app.run(host='0.0.0.0',port=int(os.getenv('PORT','5000')))
